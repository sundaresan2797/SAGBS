"""Downloading the Attendance Summary as an Excel file - it must show exactly what the screen shows."""

from datetime import date
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import update

from app.models.employee import Employee
from app.models.payroll import SalaryRecord
from app.models.user import User

YEAR, MONTH = 2026, 3


async def _login(client: AsyncClient, user: User, password: str) -> dict[str, str]:
    response = await client.post("/api/v1/auth/login", json={"email": user.email, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['data']['access_token']}"}


@pytest.mark.asyncio
async def test_export_matches_the_screen_exactly(
    client: AsyncClient, db_session, manager_with_team: dict, accountant_user: User, demo_password: str
):
    report = manager_with_team["report_employee"]
    await db_session.execute(update(Employee).values(date_of_joining=date(2020, 1, 1)))
    db_session.add(SalaryRecord(employee_id=report.id, basic_pay=24000, allowances=4000, deductions=1500,
                                effective_from=date(2025, 1, 1)))
    await db_session.flush()

    headers = await _login(client, accountant_user, demo_password)
    screen = (await client.get("/api/v1/payroll/attendance-summary", params={"year": YEAR, "month": MONTH},
                               headers=headers)).json()["data"]
    row = next(r for r in screen["items"] if r["employee_id"] == report.id)

    response = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                                headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "March_2026" in response.headers["content-disposition"]

    sheet = load_workbook(BytesIO(response.content)).active
    header = [cell.value for cell in sheet[4]]
    assert header[:4] == ["Employee code", "Employee name", "Total days", "Working days"]

    by_code = {r[0].value: r for r in sheet.iter_rows(min_row=5) if r[0].value == report.employee_code}
    excel_row = by_code[report.employee_code]
    values = {header[i]: excel_row[i].value for i in range(len(header))}
    assert values["Employee name"] == report.full_name
    assert values["Total days"] == row["total_days"]
    assert values["Present"] == row["present"]
    assert values["Absent"] == row["absent"]
    assert values["Payable days"] == row["payable_days"]
    assert float(values["Monthly salary (Rs.)"]) == row["monthly_salary"]
    assert values["Pay by attendance (Rs.)"] == row["pay_for_month"]


@pytest.mark.asyncio
async def test_export_requires_payroll_view(
    client: AsyncClient, manager_with_team: dict, hr_user: User, accountant_user: User, demo_password: str
):
    denied = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                              headers=await _login(client, manager_with_team["manager_user"], demo_password))
    assert denied.status_code == 403

    denied_hr = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                                 headers=await _login(client, hr_user, demo_password))
    assert denied_hr.status_code == 403

    allowed = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                               headers=await _login(client, accountant_user, demo_password))
    assert allowed.status_code == 200

    anonymous = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH})
    assert anonymous.status_code in (401, 403)


@pytest.mark.asyncio
async def test_export_with_nobody_paid_still_produces_a_valid_file(
    client: AsyncClient, accountant_user: User, demo_password: str
):
    """No salary structures at all for that month - the file must still open cleanly, just with no rows."""
    headers = await _login(client, accountant_user, demo_password)
    response = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": 2019, "month": 1},
                                headers=headers)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet["A1"].value == "Attendance Summary - January 2019"
    assert "No one has a salary" in sheet["A5"].value


@pytest.mark.asyncio
async def test_export_footer_shows_who_and_when(
    client: AsyncClient, accountant_user: User, demo_password: str
):
    headers = await _login(client, accountant_user, demo_password)
    response = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                                headers=headers)
    sheet = load_workbook(BytesIO(response.content)).active
    assert accountant_user.email in sheet["A2"].value
    values = [sheet.cell(row=r, column=1).value for r in range(1, 20)]
    assert any(v and "Before you pay" in v for v in values), values


@pytest.mark.asyncio
async def test_downloads_expose_their_filename_across_origins(
    client: AsyncClient, accountant_user: User, demo_password: str
):
    """Content-Disposition must be readable by browser JS from a different origin (a separate frontend port,
    like this app's own dev setup) - otherwise every download silently falls back to a generic filename."""
    headers = await _login(client, accountant_user, demo_password)
    headers["Origin"] = "http://localhost:5173"  # a cross-origin request is what actually triggers CORS headers
    response = await client.get("/api/v1/payroll/attendance-summary/export", params={"year": YEAR, "month": MONTH},
                                headers=headers)
    assert "content-disposition" in response.headers.get("access-control-expose-headers", "").lower()
