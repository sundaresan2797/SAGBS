"""Excel import of the first real data. The promises being pinned down:

- the CHECK finds every mistake, with the row number, and changes nothing;
- the LOAD is all-or-nothing;
- what is loaded is exactly what a person would have typed in (managers, roles, salaries, leave, audit trail,
  working temporary logins).
"""

from datetime import date
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker

from app.core.security import hash_password
from app.models import Department, Designation, Employee, LeaveBalance, LeaveType, SalaryRecord, User
from app.models.audit_log import AuditLog
from app.services import import_service, payroll_service
from tests.conftest import seed_roles_permissions
from app.services.import_service import (
    EMPLOYEE_COLUMNS,
    LEAVE_COLUMNS,
    LEAVE_SHEET,
    PEOPLE_SHEET,
    build_plan,
    execute_plan,
    write_login_sheet,
    write_template,
)


def person(email, first, last="Person", **kw):
    row = {
        "email": email, "first_name": first, "last_name": last, "date_of_joining": date(2022, 4, 1),
        "department": "Operations", "job_title": "Operations Executive",
    }
    row.update(kw)
    return row


def make_file(tmp_path: Path, people: list[dict], leave: list[dict] | None = None, name="people.xlsx") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = PEOPLE_SHEET
    sheet.append([c[0] for c in EMPLOYEE_COLUMNS])
    for row in people:
        sheet.append([row.get(c[0]) for c in EMPLOYEE_COLUMNS])
    if leave is not None:
        leave_sheet = workbook.create_sheet(LEAVE_SHEET)
        leave_sheet.append([c[0] for c in LEAVE_COLUMNS])
        for row in leave:
            leave_sheet.append([row.get(c[0]) for c in LEAVE_COLUMNS])
    path = tmp_path / name
    workbook.save(path)
    return path


@pytest.fixture
async def base(db_session, roles, admin_user):
    """Roles, an Admin login and the leave types, all committed so a rollback in a test cannot remove them."""
    for code, name, paid, days in (("CASUAL", "Casual Leave", True, 12), ("SICK", "Sick Leave", True, 10), ("UNPAID", "Unpaid Leave", False, 0)):
        db_session.add(LeaveType(code=code, name=name, is_paid=paid, default_annual_days=days, is_active=True))
    await db_session.commit()
    return admin_user


def messages(plan) -> str:
    return "\n".join(str(p) for p in plan.errors)


TEAM = [
    # deliberately out of order: an employee listed BEFORE their manager
    person("asha@sat.com", "Asha", "Nair", reports_to_email="mani@sat.com", basic_pay=30000, allowances=2000, deductions=500,
           job_title="Dispatcher"),
    person("mani@sat.com", "Mani", "Kumar", role="manager", reports_to_email="founder@sat.com", basic_pay=60000, allowances=5000,
           salary_from="01/05/2022", job_title="Operations Manager"),
    person("founder@sat.com", "Sundar", "S", role="FOUNDER", department="Management", job_title="Founder", employment_type="full time"),
    person("ravi@sat.com", "Ravi", "K", reports_to_email="mani@sat.com", phone="9876543210", date_of_birth="14/03/1990",
           work_location="Head Office", job_title="Dispatcher"),
]


@pytest.mark.asyncio
async def test_a_good_file_loads_exactly_as_typed_in_by_hand(client: AsyncClient, db_session, base, tmp_path):
    leave = [
        {"email": "asha@sat.com", "leave_type_code": "casual", "days_used": 2.5},
        {"email": "mani@sat.com", "leave_type_code": "SICK", "days_used": 1},
    ]
    plan = await build_plan(db_session, make_file(tmp_path, TEAM, leave))
    assert plan.ok, messages(plan)
    assert [p.email for p in plan.people][0] == "founder@sat.com"  # managers are created before their teams
    assert plan.new_departments == ["Operations", "Management"] or set(plan.new_departments) == {"Operations", "Management"}
    assert len(plan.new_job_titles) == 3 and len(plan.leave_rows) == 2  # Dispatcher is shared by two people

    result = await execute_plan(db_session, plan, base)
    assert (result.employees, result.departments, result.job_titles, result.salaries, result.leave_balances) == (4, 2, 3, 2, 2)

    people = {e.email: e for e in (await db_session.execute(select(Employee))).scalars()}
    assert set(people) == {"founder@sat.com", "mani@sat.com", "asha@sat.com", "ravi@sat.com"}
    assert people["founder@sat.com"].manager_id is None
    assert people["mani@sat.com"].manager_id == people["founder@sat.com"].id
    assert people["asha@sat.com"].manager_id == people["mani@sat.com"].id
    assert people["ravi@sat.com"].manager_id == people["mani@sat.com"].id
    assert sorted(e.employee_code for e in people.values()) == ["EMP-0001", "EMP-0002", "EMP-0003", "EMP-0004"]
    ravi = people["ravi@sat.com"]
    assert (ravi.phone, ravi.date_of_birth, ravi.work_location) == ("9876543210", date(1990, 3, 14), "Head Office")
    assert people["asha@sat.com"].date_of_joining == date(2022, 4, 1)

    # roles ("manager" typed in lower case, "full time" with a space) and departments / job titles
    roles = {u.email: u.role.name for u in (await db_session.execute(select(User))).scalars()}
    assert roles["founder@sat.com"] == "FOUNDER" and roles["mani@sat.com"] == "MANAGER" and roles["asha@sat.com"] == "EMPLOYEE"
    assert people["founder@sat.com"].employment_type.value == "FULL_TIME"
    departments = {d.name for d in (await db_session.execute(select(Department))).scalars()}
    assert departments == {"Operations", "Management"}
    dispatcher = [t for t in (await db_session.execute(select(Designation))).scalars() if t.name == "Dispatcher"]
    assert len(dispatcher) == 1  # one job title, shared by Asha and Ravi, not two

    # salaries: default start date is the joining date; an explicit one is kept
    salaries = {s.employee_id: s for s in (await db_session.execute(select(SalaryRecord))).scalars()}
    assert len(salaries) == 2
    asha, mani = salaries[people["asha@sat.com"].id], salaries[people["mani@sat.com"].id]
    assert (float(asha.basic_pay), float(asha.allowances), float(asha.deductions), asha.effective_from) == (30000, 2000, 500, date(2022, 4, 1))
    assert (float(mani.basic_pay), mani.effective_from) == (60000, date(2022, 5, 1))
    assert ravi.id not in salaries  # no salary given -> none created

    # leave already taken: used days as given, allocation = the type's yearly quota
    balances = {(b.employee_id, b.leave_type.code): b for b in (await db_session.execute(select(LeaveBalance))).scalars()}
    casual = balances[(people["asha@sat.com"].id, "CASUAL")]
    assert (float(casual.used_days), float(casual.allocated_days), casual.year) == (2.5, 12.0, date.today().year)
    assert float(balances[(people["mani@sat.com"].id, "SICK")].used_days) == 1.0

    # audit trail: one entry per employee, plus the summary
    audit = [a.action for a in (await db_session.execute(select(AuditLog))).scalars()]
    assert audit.count("CREATE_EMPLOYEE") == 4 and audit.count("IMPORT_DATA") == 1
    assert audit.count("CREATE_DEPARTMENT") == 2

    # every temporary password really works, and forces a change at first sign-in
    assert len(result.logins) == 4 and len({l.temporary_password for l in result.logins}) == 4
    for login in result.logins:
        response = await client.post("/api/v1/auth/login", json={"email": login.email, "password": login.temporary_password})
        assert response.status_code == 200, f"{login.email}: {response.text}"
        assert response.json()["data"]["user"]["must_change_password"] is True

    # the private sheet lists them, and only the owner can read it
    sheet = write_login_sheet(result, tmp_path / "private")
    assert oct(sheet.stat().st_mode)[-3:] == "600"
    text = sheet.read_text()
    assert "asha@sat.com" in text and result.logins[0].temporary_password in text


@pytest.mark.asyncio
async def test_the_leave_balance_the_import_creates_is_what_the_leave_rules_use(client: AsyncClient, db_session, base, tmp_path):
    """Casual leave: 12 a year, 10 already taken. The person can apply for 2 days but not 3."""
    plan = await build_plan(db_session, make_file(tmp_path, TEAM, [{"email": "asha@sat.com", "leave_type_code": "CASUAL", "days_used": 10}]))
    result = await execute_plan(db_session, plan, base)
    login = next(l for l in result.logins if l.email == "asha@sat.com")
    token = (await client.post("/api/v1/auth/login", json={"email": login.email, "password": login.temporary_password})).json()["data"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    balance = await client.get("/api/v1/leaves/balance", headers=headers)
    casual = next(b for b in balance.json()["data"] if b["leave_type_code"] == "CASUAL")
    assert (casual["allocated_days"], casual["used_days"], casual["remaining_days"]) == (12.0, 10.0, 2.0)


@pytest.mark.asyncio
async def test_the_check_finds_every_mistake_with_its_row_number_and_changes_nothing(db_session, base, tmp_path):
    bad = [
        person("dup@sat.com", "A"),                                                              # row 2 (fine, first use)
        person("dup@sat.com", "B"),                                                              # row 3 duplicate
        person("", "NoEmail"),                                                                   # row 4
        person("not-an-email", "Bad"),                                                           # row 5
        person("future@sat.com", "F", date_of_joining=date(2999, 1, 1)),                         # row 6
        person("baddate@sat.com", "D", date_of_joining="31/31/2020"),                            # row 7
        person("noboss@sat.com", "N", reports_to_email="ghost@sat.com"),                         # row 8 unknown manager
        person("self@sat.com", "S", reports_to_email="self@sat.com"),                            # row 9
        person("role@sat.com", "R", role="CEO"),                                                 # row 10
        person("admin@sat.com", "Ad", role="ADMIN"),                                             # row 11
        person("type@sat.com", "T", employment_type="SOMETIMES"),                                # row 12
        person("neg@sat.com", "Ng", basic_pay=-5),                                               # row 13
        person("ded@sat.com", "Dd", basic_pay=1000, allowances=0, deductions=5000),              # row 14
        person("word@sat.com", "W", basic_pay="lots"),                                           # row 15
        person("example-x@sat.com", "E"),                                                        # row 16 (not the template marker)
        person("EXAMPLE-delete.this.row@company.com", "Example"),                                # row 17 template example
        person(base.email, "Existing"),                                                          # row 18 already in the system
        person("nolast@sat.com", "L", last_name=None),                                           # row 19
        person("nodept@sat.com", "M", department=None),                                          # row 20
    ]
    plan = await build_plan(db_session, make_file(tmp_path, bad))
    assert not plan.ok
    text = messages(plan)
    expected = {
        3: "already used in row 2", 4: "email is missing", 5: "not a valid email", 6: "in the future", 7: "not a date",
        8: "ghost@sat.com is not in this file", 9: "report to themselves", 10: "'CEO' is not valid", 11: "ADMIN login is created by the setup",
        12: "employment type", 13: "negative", 14: "deductions are more than", 15: "not a number", 17: "example row",
        18: "already has a login", 19: "last name is missing", 20: "department is missing",
    }
    for row, fragment in expected.items():
        assert any(p.line == row and fragment in p.message for p in plan.errors), f"row {row}: {fragment}\n{text}"
    assert not any(p.line == 16 for p in plan.errors)  # only the exact template marker is refused

    # nothing was created, and a plan with errors refuses to load
    assert (await db_session.execute(select(func.count()).select_from(Employee))).scalar_one() == 0
    assert (await db_session.execute(select(func.count()).select_from(Department))).scalar_one() == 0
    with pytest.raises(ValueError):
        await execute_plan(db_session, plan, base)


@pytest.mark.asyncio
async def test_people_who_report_to_each_other_in_a_circle_are_caught(db_session, base, tmp_path):
    plan = await build_plan(db_session, make_file(tmp_path, [
        person("a@sat.com", "A", reports_to_email="b@sat.com"), person("b@sat.com", "B", reports_to_email="c@sat.com"),
        person("c@sat.com", "C", reports_to_email="a@sat.com"),
    ]))
    assert not plan.ok and "circle" in messages(plan)


@pytest.mark.asyncio
async def test_leave_sheet_mistakes_are_caught(db_session, base, tmp_path):
    leave = [
        {"email": "asha@sat.com", "leave_type_code": "NOPE", "days_used": 1},          # row 2
        {"email": "nobody@sat.com", "leave_type_code": "CASUAL", "days_used": 1},      # row 3
        {"email": "asha@sat.com", "leave_type_code": "CASUAL", "days_used": 13},       # row 4 over the quota of 12
        {"email": "asha@sat.com", "leave_type_code": "SICK", "days_used": 1.3},        # row 5 not a half-day step
        {"email": "asha@sat.com", "leave_type_code": "SICK", "days_used": -1},         # row 6
        {"email": "asha@sat.com", "leave_type_code": "SICK", "days_used": 2},          # row 7 (SICK listed again)
        {"email": "asha@sat.com", "leave_type_code": "CASUAL", "days_used": None},     # row 8
    ]
    plan = await build_plan(db_session, make_file(tmp_path, TEAM, leave))
    text = messages(plan)
    for row, fragment in {2: "does not exist", 3: "not on the People sheet", 4: "more than the yearly", 5: "whole or half days",
                          6: "whole or half days", 7: "listed twice", 8: "missing"}.items():
        assert any(p.line == row and p.sheet == LEAVE_SHEET and fragment in p.message for p in plan.errors), f"row {row}: {fragment}\n{text}"


@pytest.mark.asyncio
async def test_warnings_flag_things_that_are_legal_but_usually_wrong(db_session, base, tmp_path):
    people = [
        person("boss@sat.com", "Boss", role="FOUNDER"),
        person("lead@sat.com", "Lead", reports_to_email="boss@sat.com", basic_pay=1000),                # has a team but role EMPLOYEE
        person("staff@sat.com", "Staff", reports_to_email="lead@sat.com", basic_pay=1000),
        person("loose@sat.com", "Loose", basic_pay=1000),                                               # no manager, not the founder
    ]
    plan = await build_plan(db_session, make_file(tmp_path, people))
    assert plan.ok
    text = "\n".join(str(w) for w in plan.warnings)
    assert "Lead Person has people reporting to them but the role is EMPLOYEE" in text
    assert "Loose Person has no manager" in text
    assert "Boss Person has no salary" in text and "Lead Person has no salary" not in text

    two_founders = await build_plan(db_session, make_file(tmp_path, [person("f1@sat.com", "A", role="FOUNDER"), person("f2@sat.com", "B", role="FOUNDER")], name="two.xlsx"))
    assert "2 people with the role FOUNDER" in "\n".join(str(w) for w in two_founders.warnings)


@pytest.mark.asyncio
async def test_a_failure_halfway_leaves_nothing_behind(test_engine, tmp_path, monkeypatch):
    """Uses real, separate database sessions (not the shared test session), so the rollback is the real one."""
    factory = async_sessionmaker(test_engine, expire_on_commit=False)
    async with factory() as seed:
        roles = await seed_roles_permissions(seed)
        seed.add(User(email="admin@sat.com", hashed_password=hash_password("Whatever-123"), role_id=roles["ADMIN"].id, is_active=True))
        seed.add(LeaveType(code="CASUAL", name="Casual Leave", is_paid=True, default_annual_days=12, is_active=True))
        await seed.commit()

    calls = {"n": 0}
    real = payroll_service.set_salary_record

    async def flaky(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("the second salary fails")
        return await real(*args, **kwargs)

    path = make_file(tmp_path, TEAM, [{"email": "asha@sat.com", "leave_type_code": "CASUAL", "days_used": 2}])
    monkeypatch.setattr(payroll_service, "set_salary_record", flaky)
    async with factory() as session:
        admin = (await session.execute(select(User))).scalar_one()
        plan = await build_plan(session, path)
        assert plan.ok
        with pytest.raises(RuntimeError, match="second salary"):
            await execute_plan(session, plan, admin)
        assert "commit" not in vars(session)  # the session is back to normal

    # ALL of it was undone: no departments, job titles, employees, salaries, leave, logins or audit entries
    async with factory() as check:
        for model in (Department, Designation, Employee, SalaryRecord, LeaveBalance):
            assert (await check.execute(select(func.count()).select_from(model))).scalar_one() == 0, model.__name__
        assert (await check.execute(select(func.count()).select_from(User))).scalar_one() == 1  # only the Admin
        assert (await check.execute(select(func.count()).select_from(AuditLog))).scalar_one() == 0

    # ...and once the problem is gone the very same file loads completely
    monkeypatch.setattr(payroll_service, "set_salary_record", real)
    async with factory() as session:
        admin = (await session.execute(select(User))).scalar_one()
        result = await execute_plan(session, await build_plan(session, path), admin)
    assert (result.employees, result.leave_balances) == (4, 1)
    async with factory() as check:
        assert (await check.execute(select(func.count()).select_from(Employee))).scalar_one() == 4


@pytest.mark.asyncio
async def test_a_second_load_of_the_same_file_is_refused_not_duplicated(db_session, base, tmp_path):
    path = make_file(tmp_path, TEAM)
    await execute_plan(db_session, await build_plan(db_session, path), base)
    again = await build_plan(db_session, path)
    assert not again.ok and sum("already has a login" in p.message for p in again.errors) == 4
    assert (await db_session.execute(select(func.count()).select_from(Employee))).scalar_one() == 4


@pytest.mark.asyncio
async def test_people_can_be_added_under_managers_who_already_exist(db_session, base, tmp_path):
    await execute_plan(db_session, await build_plan(db_session, make_file(tmp_path, TEAM)), base)
    later = make_file(tmp_path, [person("new@sat.com", "Newbie", reports_to_email="mani@sat.com", basic_pay=20000, job_title="Dispatcher")], name="later.xlsx")
    plan = await build_plan(db_session, later)
    assert plan.ok, messages(plan)
    assert plan.new_departments == [] and plan.new_job_titles == []  # Operations / Dispatcher already exist
    await execute_plan(db_session, plan, base)
    new = (await db_session.execute(select(Employee).where(Employee.email == "new@sat.com"))).scalar_one()
    mani = (await db_session.execute(select(Employee).where(Employee.email == "mani@sat.com"))).scalar_one()
    assert new.manager_id == mani.id


@pytest.mark.asyncio
async def test_unreadable_or_incomplete_files_give_a_clear_message(db_session, base, tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_text("this is not excel")
    assert "not a readable Excel" in messages(await build_plan(db_session, junk))

    workbook = Workbook()
    workbook.active.title = "Sheet1"
    other = tmp_path / "wrongsheet.xlsx"
    workbook.save(other)
    assert "no sheet named 'People'" in messages(await build_plan(db_session, other))

    workbook = Workbook()
    workbook.active.title = PEOPLE_SHEET
    workbook.active.append(["email", "first_name"])  # required columns missing
    short = tmp_path / "short.xlsx"
    workbook.save(short)
    assert "columns are missing" in messages(await build_plan(db_session, short))

    empty = make_file(tmp_path, [], name="empty.xlsx")
    assert "no employees in the file" in messages(await build_plan(db_session, empty))


@pytest.mark.asyncio
async def test_the_template_is_ready_to_use_and_its_example_row_cannot_be_loaded_by_accident(db_session, base, tmp_path):
    template = tmp_path / "template.xlsx"
    write_template(template)
    workbook = load_workbook(template)
    assert set(workbook.sheetnames) == {"Read me first", PEOPLE_SHEET, LEAVE_SHEET}
    assert [c.value for c in workbook[PEOPLE_SHEET][1]] == [c[0] for c in EMPLOYEE_COLUMNS]
    plan = await build_plan(db_session, template)
    assert not plan.ok and "example row" in messages(plan)  # forgetting to delete it is caught, not loaded

    # a template with the example row and its leave line removed reports "no employees", not a crash
    workbook[PEOPLE_SHEET].delete_rows(2)
    workbook[LEAVE_SHEET].delete_rows(2)
    cleaned = tmp_path / "cleaned.xlsx"
    workbook.save(cleaned)
    assert "no employees in the file" in messages(await build_plan(db_session, cleaned))
